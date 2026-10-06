"""把 db-dictionary.json 渲染成一份带样式的 Excel 工作簿。"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

src = Path(sys.argv[1])
dst = Path(sys.argv[2])
data = json.loads(src.read_text(encoding="utf-8"))

HEAD_FILL = PatternFill("solid", fgColor="1F4E79")
HEAD_FONT = Font(color="FFFFFF", bold=True, size=11)
ALT_FILL = PatternFill("solid", fgColor="F2F7FC")
WHITE_FILL = PatternFill("solid", fgColor="FFFFFF")
BORDER = Border(*[Side(style="thin", color="D0D7E2")] * 4)
TOP_WRAP = Alignment(vertical="top", wrap_text=True)
TOP = Alignment(vertical="top")

wb = Workbook()

# ---- Sheet 1: 表清单 ----
ws = wb.active
ws.title = "表清单"
headers = ["序号", "表名", "表说明", "字段数", "估算行数", "引擎", "业务白名单"]
ws.append(headers)
for idx, t in enumerate(data["tables"], 1):
    ws.append([idx, t["name"], t["comment"], len(t["columns"]), t["rows"], t["engine"], "是" if t["whitelisted"] else ""])

# ---- Sheet 2: 字段明细 ----
ws2 = wb.create_sheet("字段明细")
headers2 = ["表名", "表说明", "字段名", "字段类型", "可空", "键", "默认值", "字段说明"]
ws2.append(headers2)
for t in data["tables"]:
    for c in t["columns"]:
        default = "" if c["default"] is None else str(c["default"])
        ws2.append([t["name"], t["comment"], c["name"], c["type"],
                    "YES" if c["nullable"] else "NO", c["key"], default, c["comment"]])

def style(sheet, widths, freeze="A2"):
    for cell in sheet[1]:
        cell.fill = HEAD_FILL
        cell.font = HEAD_FONT
        cell.alignment = Alignment(vertical="center", horizontal="center")
        cell.border = BORDER
    sheet.row_dimensions[1].height = 22
    for i, w in enumerate(widths, 1):
        sheet.column_dimensions[get_column_letter(i)].width = w
    for r, row in enumerate(sheet.iter_rows(min_row=2), start=2):
        fill = ALT_FILL if r % 2 == 0 else WHITE_FILL
        for cell in row:
            cell.fill = fill
            cell.border = BORDER
            cell.alignment = TOP_WRAP
    sheet.freeze_panes = freeze
    sheet.auto_filter.ref = sheet.dimensions

style(ws, [6, 34, 46, 9, 11, 10, 12])
style(ws2, [30, 34, 26, 26, 7, 8, 16, 52])

# ---- Sheet 3: 概览 ----
ws3 = wb.create_sheet("概览")
rows = [
    ["数据库", data["connection"]["database"]],
    ["主机", f'{data["connection"]["host"]}:{data["connection"]["port"]}'],
    ["账号", data["connection"]["username"]],
    ["导出时间", data["generatedAt"]],
    ["基础表数量", data["totalTables"]],
    ["视图数量", data["totalViews"]],
    ["字段总数", data["totalColumns"]],
    ["无表备注的表", sum(1 for t in data["tables"] if t["comment"] == "未填写表备注")],
    ["无字段备注的字段", sum(1 for t in data["tables"] for c in t["columns"] if not c["comment"])],
    ["配置白名单", "、".join(data.get("allowedTables", []))],
]
ws3.append(["项目", "内容"])
for r in rows:
    ws3.append(r)
style(ws3, [22, 72], freeze="A2")
for row in ws3.iter_rows(min_row=2):
    row[0].font = Font(bold=True)

wb.save(dst)
print(json.dumps({"ok": True, "path": str(dst)}, ensure_ascii=False))
